from pathlib import Path
from datetime import datetime
from collections import Counter
import json
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from PIL import Image

root = Path(__file__).resolve().parent.parent
book = root.parent / '01a0a8ac-ffe5-7202-b9f0-17032fb51354' / 'client-requirements.xlsx'
seed = json.loads((root / 'mock_seed.json').read_text(encoding='utf-8'))
users = {u['id']: u for u in seed['users']}
teams = {t['id']: t for t in seed['teams']}
tasks = {t['id']: t for t in seed['tasks']}
assert len(users) == 9 and len({u['phoneE164'] for u in users.values()}) == 9
assert Counter(u['role'] for u in users.values()) == {'employee': 5, 'admin': 3, 'superAdmin': 1}
for user in users.values():
    assert user['managerId'] is None or user['managerId'] in users
    assert user['teamId'] is None or user['teamId'] in teams
for team in teams.values():
    manager = users[team['managerId']]
    assert manager['role'] == 'admin' and manager['managedTeamId'] == team['id']
for task in tasks.values():
    assert task['assigneeId'] in users and task['createdById'] in users
    assert users[task['assigneeId']]['teamId'] == task['teamId']
    assert 0 <= task['progress'] <= 100
    if task['status'] in ('done', 'inReview'):
        assert task['progress'] == 100
    assert (task['status'] == 'done') == bool(task['completedAt'])
parse = lambda value: datetime.fromisoformat(value.replace('Z', '+00:00'))
hours = Counter()
for entry in seed['workEntries']:
    assert tasks[entry['taskId']]['assigneeId'] == entry['employeeId']
    start, end = parse(entry['startedAt']), parse(entry['endedAt'])
    assert start < end <= parse(seed['referenceNow'])
    hours[entry['employeeId']] += (end-start).total_seconds()/3600
for employee in hours:
    entries = sorted([e for e in seed['workEntries'] if e['employeeId'] == employee], key=lambda e:e['startedAt'])
    assert all(parse(a['endedAt']) <= parse(b['startedAt']) for a,b in zip(entries,entries[1:]))
assert sum(hours.values()) == 54 and hours['E001'] == 10.5
assert sum(t['status'] != 'done' and parse(t['dueAt']) < parse(seed['referenceNow']) for t in tasks.values()) == 2
for collection in ('taskUpdates', 'activity'):
    for record in seed[collection]:
        assert record['actorId'] in users
        assert record.get('taskId',record.get('entityId')) in tasks

screens = sorted((root/'ui').glob('[0-9][0-9]-*.svg'))
assert len(screens) == 17
for svg in screens:
    node=ET.parse(svg).getroot()
    assert node.attrib['viewBox'] == '0 0 390 844'
    assert not node.findall('.//{http://www.w3.org/2000/svg}script')
    assert not node.findall('.//{http://www.w3.org/2000/svg}image')
ET.parse(root/'ui/teamflow-ui-board.svg')

wb = openpyxl.load_workbook(book, data_only=False)
assert wb.sheetnames == ['Requirements','Role access','Demo logins','API catalogue','Workshop schedule']
for i,user in enumerate(seed['users'],start=5):
    cell=wb['Demo logins'][f'E{i}']
    assert cell.data_type == 's'
    assert cell.value.replace(' ','').replace('-','') == user['phoneE164']
assert len(wb['Requirements'].tables) == 1 and wb['Requirements'].freeze_panes == 'A5'
cached=openpyxl.load_workbook(book,data_only=True)
assert cached['Workshop schedule']['B20'].value == 360
assert cached['Workshop schedule']['C20'].value == 6
assert cached['Workshop schedule']['B21'].value == 360
assert cached['Workshop schedule']['C21'].value == 6

req=Image.open(root/'workbook-qa/Requirements.png')
req.crop((0,1080,min(700,req.width),min(1760,req.height))).save(root/'workbook-qa/Requirements-detail.png')
files = [root/'README.md',root/'IMPLEMENTATION_PLAN.md',root/'mock_seed.json',*screens,
         root/'ui/teamflow-ui-board.svg',root/'ui/README.md',root/'ui/index.html',root/'ui/role-dashboards.png']
zip_path=root/'teamflow-workshop-pack.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as archive:
    for file in files:
        archive.write(file,file.relative_to(root).as_posix())
    archive.write(book,'client-requirements.xlsx')
with zipfile.ZipFile(zip_path) as archive:
    assert archive.testzip() is None
    assert 'client-requirements.xlsx' in archive.namelist()
print(json.dumps({'accounts':len(users),'tasks':len(tasks),'work_entries':len(seed['workEntries']),
 'logged_hours':sum(hours.values()),'svg_screens':len(screens),'workbook_sheets':len(wb.sheetnames),
 'teaching_hours':12,'archive_files':len(files)+1,'zip_bytes':zip_path.stat().st_size}))
